"""The questions workbook: what the sheet does not say, one row per item, with empty columns to fill in."""
import json, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
r = json.load(open(sys.argv[1], encoding='utf-8'))
wb = openpyxl.Workbook()
head = PatternFill('solid', fgColor='1D5FBF'); fill = PatternFill('solid', fgColor='FFF7E0'); bold = Font(bold=True, color='FFFFFF')

def sheet(ws, cols, rows, widths, fill_from):
    ws.append(cols)
    for c in ws[1]: c.fill = head; c.font = bold; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in rows: ws.append(row)
    for i, w in enumerate(widths, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[fill_from:]: c.fill = fill
    ws.freeze_panes = 'C2'

cos = {c[0]: c for c in r['companies']}
ws = wb.active; ws.title = 'Companies'
sheet(ws, ['Code (in the script)', 'Arabic legal name (from the sheet)', 'English name used', 'Correct English name', 'Correct code', 'Commercial reg. no.', 'Address', 'Telephone', 'Email'],
      [[c[0], c[2], c[1], None, None, None, None, None, None] for c in r['companies']], [12, 55, 38, 38, 12, 18, 30, 14, 26], 3)

ws = wb.create_sheet('Employees')
cols = ['Company', 'Emp Code', 'Code made up?', 'Name', 'Works at (if not the company)', 'Nationality', 'Joining date', 'Marital status',
        'Mobile', 'Department', 'Designation used (work permit)', 'Correct designation', 'Allowances (KWD)', 'Bank', 'IBAN',
        'Work permit no.', 'Work permit expiry', 'Passport ✓ / X means?', 'Sheet notes']
rows = []
for e in r['emps']:
    rows.append([e['company'], e['code'], 'YES - confirm' if e['generated'] else None,
                 ' '.join(n for n in (e['first'], e['middle'], e['last']) if n), e['workplace'],
                 e['nat'], None, None, None, None, e['job_en'], None, None, None, None, None, None, e['custody'], e['notes']])
sheet(ws, cols, rows, [9, 9, 13, 36, 14, 13, 13, 13, 14, 16, 30, 24, 12, 12, 26, 15, 15, 14, 40], 5)
for row in ws.iter_rows(min_row=2):
    if row[5].value: row[5].fill = PatternFill()          # nationality known
    row[10].fill = PatternFill(); row[17].fill = PatternFill(); row[18].fill = PatternFill()

ws = wb.create_sheet('Vehicles')
sheet(ws, ['Company', 'Type', 'Brand', 'Model', 'Plate', 'Registration expiry', 'Sheet description / holder', 'Assigned employee (code)', 'Brand (if blank)'],
      [[v[0], v[1], v[2], v[3], v[5], v[6], v[7], None, None] for v in r['vehicles']], [9, 16, 14, 18, 12, 14, 60, 18, 14], 7)
wb.save(sys.argv[2])
print('saved')
